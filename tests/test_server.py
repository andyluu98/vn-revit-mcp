import asyncio
import json

import openpyxl
import pytest

from vn_revit_mcp import excel, tong_hop
from vn_revit_mcp.server import mcp

EXPECTED_TOOLS = {
    "kiem_tra_ket_noi", "thong_tin_du_an", "dem_theo_hang_muc", "thong_ke_cua", "thong_ke_phong",
    "khoi_luong_tuong", "khoi_luong_san", "khoi_luong_ket_cau", "kiem_tra_tham_so",
    "xuat_excel_khoi_luong", "dat_tham_so",
}


def call(name, **arguments):
    result = asyncio.run(mcp.call_tool(name, arguments))
    if getattr(result, "structuredContent", None) is not None:
        data = result.structuredContent
        return data.get("result", data) if set(data) == {"result"} else data
    return json.loads(result.content[0].text)


def test_tool_list_and_annotations():
    tools = {tool.name: tool for tool in asyncio.run(mcp.list_tools())}
    assert set(tools) == EXPECTED_TOOLS
    assert tools["dat_tham_so"].annotations.destructive_hint is True
    assert tools["khoi_luong_tuong"].annotations.read_only_hint is True
    assert tools["xuat_excel_khoi_luong"].annotations.read_only_hint is False
    for tool in tools.values():
        assert tool.description


def test_connection(fake_cli):
    data = call("kiem_tra_ket_noi")
    assert data["phien_revit"][0]["pid"] == 4242


def test_wall_quantities_sum_full_precision(fake_cli):
    data = call("khoi_luong_tuong")
    group = next(g for g in data["theo_loai_va_tang"] if g["type"] == "Tuong 200")
    assert group["so_luong"] == 2
    assert group["volume_m3"] == round(3.60288 + 2.88216, 2)


def test_doors_grouped_by_type_and_level(fake_cli):
    data = call("thong_ke_cua")
    door = next(g for g in data["theo_loai"] if g["kind"] == "door")
    assert door["so_luong"] == 2
    assert door["theo_tang"] == {"Tang 1": 1, "Tang 2": 1}
    assert data["khong_co_ky_hieu"] == [102]


def test_rooms_flag_unplaced(fake_cli):
    data = call("thong_ke_phong")
    assert data["chua_dat_vi_tri"] == [303]
    assert data["tong_dien_tich_m2"] == round(20.123456 + 10.5, 2)


def test_parameter_check_resolves_vietnamese_alias(fake_cli):
    data = call("kiem_tra_tham_so", hang_muc="cửa đi", tham_so="Mark")
    assert data["category"] == "OST_Doors"


def test_parameter_check_rejects_bad_category(fake_cli):
    data = call("kiem_tra_tham_so", hang_muc="DROP TABLE", tham_so="Mark")
    assert data["loi"] == "dau_vao_khong_hop_le"


def test_set_parameter_defaults_to_dry_run(fake_cli):
    data = call("dat_tham_so", ids=[5, 6], tham_so="Comments", gia_tri="Kiểm tra")
    assert data["che_do"] == "dry_run"
    assert fake_cli()[-1]["mode"] == "dry_run"


def test_set_parameter_reports_discard(fake_cli, monkeypatch):
    monkeypatch.setenv("FAKE_PYREVIT_SCENARIO", "discard")
    data = call("dat_tham_so", ids=[5], tham_so="Comments", gia_tri="x", chay_thu=False)
    assert data["quyet_dinh"] == "discarded"


def test_set_parameter_limits(fake_cli):
    assert call("dat_tham_so", ids=list(range(1, 300)), tham_so="Comments", gia_tri="x")["loi"] == "dau_vao_khong_hop_le"
    assert call("dat_tham_so", ids=[], tham_so="Comments", gia_tri="x")["loi"] == "dau_vao_khong_hop_le"


def test_revit_selector_must_be_numeric(fake_cli):
    data = call("thong_tin_du_an", revit="2025 --mode=modify")
    assert data["loi"] == "dau_vao_khong_hop_le"


def test_excel_export_numbered_and_never_overwrites(fake_cli, tmp_path):
    first = call("xuat_excel_khoi_luong", thu_muc=str(tmp_path))
    second = call("xuat_excel_khoi_luong", thu_muc=str(tmp_path))
    assert first["file"].endswith("01_khoi-luong-revit-nha-pho-mau.xlsx")
    assert second["file"].endswith("02_khoi-luong-revit-nha-pho-mau.xlsx")

    wb = openpyxl.load_workbook(first["file"])
    assert wb.sheetnames == ["KhoiLuong", "Tuong", "San", "KetCau", "Cua", "Phong", "Nguon"]
    ws = wb["KhoiLuong"]
    headers = [ws.cell(row=4, column=c).value for c in range(1, 11)]
    assert headers == excel.BOQ_HEADERS
    names = [ws.cell(row=r, column=3).value for r in range(5, ws.max_row + 1)]
    assert any(n.startswith("Tường Tuong 200") for n in names)
    assert not any("Vach kinh" in n for n in names), "curtain walls are not masonry"
    assert ws.cell(row=5, column=8).value.startswith("=IF(")
    assert ws.cell(row=5, column=7).value is None


def test_excel_never_writes_model_text_as_formula(tmp_path):
    evil = '=HYPERLINK("http://x","bam")'
    tuong = {"theo_loai_va_tang": [{"type": evil, "level": "-1", "kind": "basic", "so_luong": 1,
                                    "thickness_mm": 200.0, "length_m": 1.0, "area_m2": 3.0, "volume_m3": 0.6}]}
    path = excel.write_boq(tmp_path, {"title": evil}, tuong, {}, {}, {}, {})
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "HYPERLINK" in cell.value:
                    assert cell.data_type == "s", (ws.title, cell.coordinate)


def test_wall_total_excludes_curtain(fake_cli):
    data = call("khoi_luong_tuong")
    assert data["tong"]["area_m2"] == round(18.0144 + 14.4108, 2)
    assert data["tong_vach_kinh"]["area_m2"] == 9.0


def test_missing_volume_flagged_in_boq():
    items = excel.boq_items({"theo_loai_va_tang": [{"type": "In place", "level": "", "kind": "in_place",
                                                    "so_luong": 1, "volume_m3": None}]}, {}, {}, {})
    assert "Cần xác nhận lại" in items[0]["dien_giai"]


def test_excel_refuses_network_and_relative_paths(fake_cli):
    assert call("xuat_excel_khoi_luong", thu_muc="\\\\server\\share")["loi"] == "dau_vao_khong_hop_le"
    assert call("xuat_excel_khoi_luong", thu_muc="khoi-luong")["loi"] == "dau_vao_khong_hop_le"


def test_round_values_only_at_display():
    rows = [{"k": "a", "v": 0.004}, {"k": "a", "v": 0.004}]
    grouped = tong_hop.group_rows(rows, ("k",), ("v",))
    assert grouped[0]["v"] == pytest.approx(0.008)
    assert tong_hop.round_values(grouped)[0]["v"] == 0.01


def test_check_ids_rejects_bool_and_negative():
    with pytest.raises(ValueError):
        tong_hop.check_ids([True])
    with pytest.raises(ValueError):
        tong_hop.check_ids([-1])
    assert tong_hop.check_ids([3, 3, 4]) == [3, 4]
