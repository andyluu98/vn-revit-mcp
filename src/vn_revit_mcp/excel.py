"""Write the Revit quantities to an Excel BOQ.

Layout follows the usual Vietnamese takeoff sheet: design quantity, waste (%) left
empty for the engineer, waste quantity and purchase quantity as formulas.
Files are numbered NN_ in the target folder and never overwritten.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .tong_hop import normalize

FONT = "Times New Roman"
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NUMBER_FORMAT = "#,##0.00"
CHO_BO_SUNG = "[Chờ bổ sung]"
_NUMBERED = re.compile(r"^(\d{2})_")
_FORMULA_START = ("=", "+", "-", "@")

BOQ_HEADERS = [
    "STT", "Mã hiệu", "Tên công việc", "Đơn vị", "Diễn giải",
    "KL thiết kế", "Hao hụt (%)", "KL hao hụt", "KL mua", "Nguồn",
]


def default_folder() -> Path:
    return Path.home() / "Documents" / "vn-revit-mcp"


def check_folder(folder: str | None) -> Path:
    """Absolute local folder; created if its parent exists. No network paths."""
    if not folder:
        path = default_folder()
        path.mkdir(parents=True, exist_ok=True)
        return path
    if folder.startswith("\\\\") or folder.startswith("//"):
        raise ValueError("Không ghi ra đường dẫn mạng. Chọn một thư mục trên máy.")
    path = Path(folder).expanduser()
    if not path.is_absolute():
        raise ValueError("Thư mục phải là đường dẫn tuyệt đối, ví dụ D:\\DuAn\\KhoiLuong.")
    path = path.resolve()
    if path.exists() and not path.is_dir():
        raise ValueError("Đường dẫn đã có nhưng không phải thư mục.")
    if not path.exists():
        if not path.parent.is_dir():
            raise ValueError("Thư mục cha không tồn tại.")
        path.mkdir()
    return path


def _slug(text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", normalize(text or "")).strip("-")
    return slug[:50] or "du-an"


def next_file(folder: Path, title: str) -> Path:
    """NN_khoi-luong-revit-<du-an>.xlsx with NN = highest number in the folder + 1."""
    numbers = [int(m.group(1)) for p in folder.iterdir() if (m := _NUMBERED.match(p.name))]
    number = (max(numbers) + 1) if numbers else 1
    while True:
        path = folder / f"{number:02d}_khoi-luong-revit-{_slug(title)}.xlsx"
        if not path.exists():
            return path
        number += 1


def _style_header(ws, row: int, count: int) -> None:
    for col in range(1, count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_table(ws, headers: list[str], rows: list[list[Any]], start_row: int = 1) -> None:
    for col, header in enumerate(headers, start=1):
        ws.cell(row=start_row, column=col, value=header)
    _style_header(ws, start_row, len(headers))
    for r, values in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str) and value[:1] in _FORMULA_START:
                # Names come from the model (possibly third-party families): never formulas.
                cell.data_type = "s"
            cell.font = Font(name=FONT)
            cell.border = BORDER
            if isinstance(value, float):
                cell.number_format = NUMBER_FORMAT
    for col in range(1, len(headers) + 1):
        width = max(len(str(headers[col - 1])), *(len(str(v[col - 1])) for v in rows if len(v) >= col)) if rows else len(headers[col - 1])
        ws.column_dimensions[get_column_letter(col)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _fmt(value: Any, digits: int = 0) -> str:
    if value is None:
        return "?"
    return f"{value:,.{digits}f}"


def boq_items(tuong: dict, san: dict, ket_cau: dict, cua: dict) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for g in tuong.get("theo_loai_va_tang", []):
        if g.get("kind") == "curtain":
            continue
        items.append({
            "ten": f"Tường {g.get('type')}, {g.get('level') or 'chưa rõ tầng'}",
            "dvt": "m3", "kl": g.get("volume_m3"),
            "dien_giai": f"{g['so_luong']} đoạn, dày {_fmt(g.get('thickness_mm'))} mm, "
                         f"dài {_fmt(g.get('length_m'), 2)} m, diện tích {_fmt(g.get('area_m2'), 2)} m2",
            "nguon": "Revit Walls: HOST_VOLUME_COMPUTED",
        })
    for g in san.get("theo_loai_va_tang", []):
        items.append({
            "ten": f"Sàn {g.get('type')}, {g.get('level') or 'chưa rõ tầng'}",
            "dvt": "m3", "kl": g.get("volume_m3"),
            "dien_giai": f"{g['so_luong']} sàn, dày {_fmt(g.get('thickness_mm'))} mm, "
                         f"diện tích {_fmt(g.get('area_m2'), 2)} m2",
            "nguon": "Revit Floors: HOST_VOLUME_COMPUTED",
        })
    kind_vn = {"column": "Cột", "framing": "Dầm", "foundation": "Móng"}
    for g in ket_cau.get("theo_loai", []):
        items.append({
            "ten": f"{kind_vn.get(g.get('kind'), g.get('kind'))} {g.get('type')}, "
                   f"{g.get('material') or 'chưa gán vật liệu'}, {g.get('level') or 'chưa rõ tầng'}",
            "dvt": "m3", "kl": g.get("volume_m3"),
            "dien_giai": f"{g['so_luong']} cấu kiện" + (f", dài {_fmt(g.get('length_m'), 2)} m" if g.get("length_m") else ""),
            "nguon": "Revit: GetMaterialVolume",
        })
    for g in cua.get("theo_loai", []):
        name = "Cửa đi" if g.get("kind") == "door" else "Cửa sổ"
        items.append({
            "ten": f"{name} {g.get('family')} {g.get('type')}",
            "dvt": "bộ", "kl": float(g["so_luong"]),
            "dien_giai": f"{_fmt(g.get('width_mm'))} x {_fmt(g.get('height_mm'))} mm; "
                         + "; ".join(f"{k or 'chưa rõ tầng'}: {v}" for k, v in g.get("theo_tang", {}).items()),
            "nguon": "Revit Doors/Windows: đếm cá thể",
        })
    for item in items:
        if item["kl"] is None:
            item["dien_giai"] += "; Cần xác nhận lại: Revit không trả khối lượng"
    return items


def write_boq(folder: Path, project: dict, tuong: dict, san: dict, ket_cau: dict, cua: dict, phong: dict) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "KhoiLuong"
    title = project.get("title") or "du-an"
    ws.cell(row=1, column=1, value=f"BẢNG KHỐI LƯỢNG TỪ MÔ HÌNH REVIT: {title}").font = Font(name=FONT, bold=True, size=13)
    ws.cell(row=2, column=1, value=(
        "Khối lượng thiết kế lấy từ mô hình Revit, chưa đối chiếu bản vẽ. "
        "Hao hụt (%) để kỹ sư điền theo định mức; mã hiệu " + CHO_BO_SUNG + "."
    )).font = Font(name=FONT, italic=True)

    items = boq_items(tuong, san, ket_cau, cua)
    rows = []
    for index, item in enumerate(items, start=1):
        rows.append([index, CHO_BO_SUNG, item["ten"], item["dvt"], item["dien_giai"],
                     item["kl"], None, None, None, item["nguon"]])
    _write_table(ws, BOQ_HEADERS, rows, start_row=4)
    for r in range(5, 5 + len(rows)):
        ws.cell(row=r, column=7).fill = INPUT_FILL
        ws.cell(row=r, column=8, value=f'=IF(G{r}="","",F{r}*G{r}/100)').number_format = NUMBER_FORMAT
        ws.cell(row=r, column=9, value=f'=IF(G{r}="","",F{r}+H{r})').number_format = NUMBER_FORMAT
        for c in (8, 9):
            ws.cell(row=r, column=c).font = Font(name=FONT)
            ws.cell(row=r, column=c).border = BORDER

    detail = wb.create_sheet("Tuong")
    _write_table(detail, ["Loại tường", "Tầng", "Kiểu", "Dày (mm)", "Số đoạn", "Dài (m)", "Diện tích (m2)", "Thể tích (m3)"],
                 [[g.get("type"), g.get("level"), g.get("kind"), g.get("thickness_mm"), g["so_luong"],
                   g.get("length_m"), g.get("area_m2"), g.get("volume_m3")] for g in tuong.get("theo_loai_va_tang", [])])
    detail = wb.create_sheet("San")
    _write_table(detail, ["Loại sàn", "Tầng", "Kết cấu", "Dày (mm)", "Số sàn", "Diện tích (m2)", "Thể tích (m3)"],
                 [[g.get("type"), g.get("level"), "có" if g.get("structural") else "không", g.get("thickness_mm"),
                   g["so_luong"], g.get("area_m2"), g.get("volume_m3")] for g in san.get("theo_loai_va_tang", [])])
    detail = wb.create_sheet("KetCau")
    _write_table(detail, ["Loại cấu kiện", "Tên loại", "Vật liệu", "Tầng", "Số cấu kiện", "Dài (m)", "Thể tích (m3)"],
                 [[g.get("kind"), g.get("type"), g.get("material"), g.get("level"), g["so_luong"],
                   g.get("length_m"), g.get("volume_m3")] for g in ket_cau.get("theo_loai", [])])
    detail = wb.create_sheet("Cua")
    _write_table(detail, ["Loại", "Family", "Tên loại", "Rộng (mm)", "Cao (mm)", "Số bộ", "Theo tầng"],
                 [[g.get("kind"), g.get("family"), g.get("type"), g.get("width_mm"), g.get("height_mm"), g["so_luong"],
                   "; ".join(f"{k}: {v}" for k, v in g.get("theo_tang", {}).items())] for g in cua.get("theo_loai", [])])
    detail = wb.create_sheet("Phong")
    _write_table(detail, ["Id", "Số phòng", "Tên phòng", "Tầng", "Diện tích (m2)", "Chu vi (m)", "Thể tích (m3)", "Trạng thái"],
                 [[r.get("id"), r.get("number"), r.get("name"), r.get("level"), r.get("area_m2"), r.get("perimeter_m"),
                   r.get("volume_m3"), r.get("status")] for r in phong.get("phong", [])])

    source = wb.create_sheet("Nguon")
    notes = [
        ["Mô hình", title],
        ["Đường dẫn file Revit", project.get("path") or ""],
        ["Phiên bản Revit", project.get("revit_version") or ""],
        ["Thời điểm trích", datetime.now().strftime("%d/%m/%Y %H:%M")],
        ["Công cụ", "vn-revit-mcp qua pyRevit agent runtime (chế độ chỉ đọc)"],
        ["Đơn vị", "Revit lưu nội bộ theo feet; đã đổi sang m, m2, m3, mm"],
        ["Tường xếp chồng", f"Bỏ qua {tuong.get('bo_qua_tuong_xep_chong_cha', 0)} tường cha, đã tính các tường con"],
        ["Tường kính (curtain)", "Có trong sheet Tuong, không đưa vào bảng khối lượng xây"],
        ["Cắt bớt", "Có" if any(d.get("bi_cat_bot") for d in (tuong, san, ket_cau, cua, phong)) else "Không"],
        ["Tường in-place", f"{tuong.get('tuong_in_place', 0)} tường; thể tích lấy theo vật liệu, thiếu thì ghi Cần xác nhận lại"],
        ["Lưu ý", "Khối lượng phụ thuộc cách dựng mô hình (giao cắt, join, vật liệu). "
                  "Phần giao giữa các cấu kiện chưa join bị tính hai lần. Đối chiếu bản vẽ trước khi dùng."],
    ]
    _write_table(source, ["Mục", "Nội dung"], notes)

    wb.calculation.fullCalcOnLoad = True
    path = next_file(folder, title)
    wb.save(path)
    return path
